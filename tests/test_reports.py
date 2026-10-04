"""Asistente de reportes: preguntas en español sobre los datos de LA empresa, Excel y aprendizaje.

Los datos se siembran directo en la BD (empleados, departamentos, identificaciones con su hora,
validadores...) para preguntar sobre ellos como lo haría una empresa real. Otra empresa con sus
propios datos comprueba que el asistente jamás los ve.
"""

from datetime import UTC, datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.clock import business_day_start
from app.core.database import SessionLocal
from app.models import (
    AssistantQuery,
    Department,
    DeviceStatus,
    Employee,
    EnrollmentStatus,
    FaceEmbedding,
    FaceEnrollment,
    FaceStatus,
    User,
    UserRole,
    Validator,
    ValidatorDevice,
    ValidatorMode,
    VerificationLog,
    VerificationMethod,
)
from app.services import maintenance_service
from tests.conftest import create_company, curp_for, login, nss_for, phone_for, rfc_for

URL = "/api/reports"


def _hire(client, headers, number: str, first: str, last: str, email: str) -> int:
    body = {
        "first_name": first,
        "last_name": last,
        "birth_date": "1990-05-10",
        "employee_number": number,
        "rfc": rfc_for(number),
        "curp": curp_for(number),
        "nss": nss_for(number),
        "phone": phone_for(number),
        "email": email,
        "password": "Empleado123",
    }
    response = client.post("/api/employees", json=body, headers=headers)
    assert response.status_code == 201, response.text
    return int(response.json()["data"]["id"])


def _company_of(email: str) -> int:
    with SessionLocal() as db:
        return int(db.scalar(select(User.company_id).where(User.email == email)) or 0)


def _log(db, company_id: int, employee_id: int | None, hours_ago_today: float, *, success=True, **extra) -> None:
    moment = (business_day_start() + timedelta(hours=hours_ago_today)).astimezone(UTC)  # la BD guarda UTC
    db.add(
        VerificationLog(
            company_id=company_id,
            employee_id=employee_id,
            method=extra.pop("method", VerificationMethod.FACE),
            success=success,
            created_at=moment,
            **extra,
        )
    )


@pytest.fixture
def company(client, company_headers) -> dict:
    """Empresa con 3 empleados (uno inactivo), 2 departamentos, identificaciones de hoy y de ayer,
    un validador con un dispositivo por autorizar, un registro facial pendiente y muestras aprendidas."""
    ana = _hire(client, company_headers, "EMP-001", "Ana", "López", "ana@empresa.com")
    beto = _hire(client, company_headers, "EMP-002", "Beto", "Ruiz", "beto@empresa.com")
    caro = _hire(client, company_headers, "EMP-003", "Caro", "Mal", "caro@empresa.com")
    company_id = _company_of("admin@empresa.com")
    with SessionLocal() as db:
        production = Department(company_id=company_id, name="Producción")
        sales = Department(company_id=company_id, name="Ventas")
        db.add_all([production, sales])
        db.flush()
        for employee_id, department in ((ana, production), (beto, production), (caro, sales)):
            db.get(Employee, employee_id).department_id = department.id  # type: ignore[union-attr]
        db.get(Employee, caro).active = False  # type: ignore[union-attr]
        # Un dato viejo con forma de fórmula (la API ya no lo acepta): en Excel debe quedar como texto.
        db.get(Employee, caro).first_name = '=HYPERLINK("http://x")'  # type: ignore[union-attr]
        db.get(Employee, ana).face_status = FaceStatus.APPROVED  # type: ignore[union-attr]
        user = User(email="caseta@empresa.com", password_hash="x", role=UserRole.VALIDATOR, company_id=company_id)
        db.add(user)
        db.flush()
        validator = Validator(
            user_id=user.id, company_id=company_id, name="Caseta Norte", mode=ValidatorMode.QR_OR_FACE
        )
        db.add(validator)
        db.flush()
        db.add(
            ValidatorDevice(
                validator_id=validator.id,
                company_id=company_id,
                key_hash="h" * 64,
                public_key="k",
                name="iPad de la caseta",
                status=DeviceStatus.PENDING,
            )
        )
        # Hoy: Ana entra 8:50 y sale 17:50; Beto entra 9:20 y falla una vez; ayer: Ana a las 9:05.
        _log(db, company_id, ana, 8 + 50 / 60, user_id=user.id, score=0.91)
        _log(db, company_id, ana, 17 + 50 / 60, user_id=user.id, score=0.93)
        _log(db, company_id, beto, 9 + 20 / 60, method=VerificationMethod.QR)
        _log(db, company_id, beto, 9.25, success=False, reason="NO_MATCH")
        _log(db, company_id, ana, -24 + 9 + 5 / 60)
        db.add(
            FaceEnrollment(
                employee_id=beto,
                company_id=company_id,
                status=EnrollmentStatus.PENDING,
                quality_score=0.8,
                samples=3,
                liveness_passed=True,
            )
        )
        for learned in (False, True, True):
            db.add(
                FaceEmbedding(
                    employee_id=ana,
                    embedding_encrypted=b"x",
                    model_name="fake-model",
                    dimension=1,
                    detection_score=0.9,
                    quality_score=0.9,
                    learned=learned,
                    matches=2 if learned else 5,
                )
            )
        db.commit()
    return {"ana": ana, "beto": beto, "caro": caro, "company_id": company_id}


def _ask(client, headers, question: str, context: dict | None = None) -> dict:
    response = client.post(f"{URL}/ask", json={"question": question, "context": context}, headers=headers)
    assert response.status_code == 200, response.text
    return response.json()["data"]


def _filters(plan: dict) -> dict:
    return {f["column"]: (f["op"], f["values"]) for f in plan["filters"]}


# ---------------------------------------------------------------- preguntas en español


def test_counts_with_filters_and_typos(client, company_headers, company):
    answer = _ask(client, company_headers, "¿Cuántos empleados activos hay?")
    assert answer["plan"]["dataset"] == "employees" and answer["plan"]["mode"] == "count"
    assert _filters(answer["plan"]) == {"active": ("in", [True])}
    assert answer["answer"].startswith("Hay 2 empleados") and answer["preview"]["rows"] == [[2]]
    assert "Datos: Empleados" in answer["understood"]

    typo = _ask(client, company_headers, "empelados de ventas")  # error de dedo y departamento por nombre
    assert typo["plan"]["dataset"] == "employees" and _filters(typo["plan"]) == {"department": ("in", ["Ventas"])}
    assert typo["preview"]["total"] == 1


def test_groups_top_and_insights(client, company_headers, company):
    groups = _ask(client, company_headers, "empleados por departamento")
    assert groups["plan"]["mode"] == "groups" and groups["plan"]["group_by"] == ["department"]
    assert groups["preview"]["rows"][0] == ["Producción", 2]
    assert any("Producción: 2 (67 % del total)" in h for h in groups["highlights"])

    top = _ask(client, company_headers, "¿Quién se identificó más hoy?")
    assert top["plan"]["dataset"] == "attendance" and top["plan"]["group_by"] == ["employee"]
    assert top["plan"]["sort"] == "metric" and top["plan"]["descending"] is True and top["plan"]["limit"] == 10
    assert top["preview"]["rows"][0][0] in ("Ana López", "Beto Ruiz") and top["preview"]["rows"][0][1] == 2


def test_failed_attempts_by_reason_and_comparison_with_the_previous_period(client, company_headers, company):
    failed = _ask(client, company_headers, "identificaciones fallidas de hoy por motivo")
    assert _filters(failed["plan"]) == {"success": ("in", [False])}
    assert failed["plan"]["group_by"] == ["reason"] and failed["preview"]["rows"] == [["Rostro no coincide", 1]]
    assert failed["plan"]["period"]["label"].startswith("hoy")

    today = _ask(client, company_headers, "¿cuántas identificaciones hubo hoy?")
    assert today["preview"]["rows"] == [[4]]
    assert any("que en el periodo anterior (1)" in h for h in today["highlights"])  # ayer hubo 1


def test_workdays_entry_exit_hours_and_late_arrivals(client, company_headers, company):
    hours = _ask(client, company_headers, "horas trabajadas por empleado hoy")
    assert hours["plan"]["dataset"] == "workdays" and hours["plan"]["metric"] == "sum"
    ana = next(row for row in hours["preview"]["rows"] if row[0] == "Ana López")
    assert ana[1] == 9  # 8:50 → 17:50

    late = _ask(client, company_headers, "¿Quién llegó después de las 9 hoy?")
    assert _filters(late["plan"]) == {"entry_hour": ("gt", [9.0])}
    assert [row[1] for row in late["preview"]["rows"]] == ["Beto Ruiz"]
    assert "Hora de entrada (decimal): más de 9:00" in late["understood"]

    vague = _ask(client, company_headers, "¿quién llegó tarde hoy?")
    assert vague["plan"]["dataset"] == "workdays" and vague["plan"]["sort"] == "entry"
    assert any("a partir de qué hora" in h for h in vague["highlights"])


def test_other_datasets_by_their_words(client, company_headers, company):
    pending = _ask(client, company_headers, "registros faciales pendientes")
    assert pending["plan"]["dataset"] == "enrollments" and _filters(pending["plan"]) == {"status": ("in", ["PENDING"])}
    assert pending["preview"]["total"] == 1

    devices = _ask(client, company_headers, "dispositivos por autorizar")
    assert devices["plan"]["dataset"] == "devices" and devices["preview"]["rows"][0][2] == "Por autorizar"

    learning = _ask(client, company_headers, "aprendizaje facial por empleado")
    assert learning["plan"]["dataset"] == "face_learning"

    keys = _ask(client, company_headers, "llaves de la api activas")
    assert keys["preview"]["total"] == 0 and "Aún no hay llaves en tu empresa" in keys["answer"]


def test_follow_up_questions_keep_the_conversation(client, company_headers, company):
    first = _ask(client, company_headers, "identificaciones de hoy")
    by_department = _ask(client, company_headers, "¿y por departamento?", first["plan"])
    assert by_department["plan"]["dataset"] == "attendance" and by_department["plan"]["group_by"] == ["department"]
    assert by_department["plan"]["period"] == first["plan"]["period"]
    assert "Sigo con la consulta anterior" in by_department["understood"]

    export = _ask(client, company_headers, "expórtalo a excel", by_department["plan"])
    assert export["export"] is True and export["plan"]["group_by"] == ["department"]


def test_help_and_unknown_questions_are_always_answered(client, company_headers, company):
    help_ = _ask(client, company_headers, "¿Qué puedes hacer?")
    assert help_["plan"] is None and "2 empleados" in help_["answer"] and help_["suggestions"]

    unknown = _ask(client, company_headers, "¿cuántos ponches hubo?")
    assert unknown["plan"] is None and "«ponche»" in unknown["answer"]
    assert {a["code"] for a in unknown["alternatives"]} and unknown["query_id"]


def test_empty_answers_explain_what_there_is(client, company_headers, company):
    empty = _ask(client, company_headers, "empleados inactivos de producción")
    assert empty["preview"]["total"] == 0
    assert "Sin esas condiciones hay 3 empleados" in empty["answer"]


# ---------------------------------------------------------------- aprende de la empresa


def test_the_assistant_learns_company_vocabulary(client, company_headers, admin_headers, company):
    unknown = _ask(client, company_headers, "¿cuántos ponches hubo?")
    learned = client.post(
        f"{URL}/feedback", json={"query_id": unknown["query_id"], "dataset": "attendance"}, headers=company_headers
    ).json()["data"]
    assert learned["learned"] == ["«ponche» → Identificaciones"]
    assert learned["answer"]["plan"]["dataset"] == "attendance" and learned["answer"]["plan"]["mode"] == "count"

    again = _ask(client, company_headers, "ponches por validador")  # ya lo entiende a la primera
    assert again["plan"]["dataset"] == "attendance" and again["plan"]["group_by"] == ["validator"]

    # Otra empresa no aprende de esta (el vocabulario puede tener nombres de su gente).
    assert create_company(client, admin_headers).status_code == 201
    other = login(client, "admin@panificadora.com", "Empresa1234")
    assert _ask(client, other, "¿cuántos ponches hubo?")["plan"] is None


def test_marking_an_answer_useless_weakens_what_it_learned(client, company_headers, company):
    unknown = _ask(client, company_headers, "¿cuántos ponches hubo?")
    client.post(
        f"{URL}/feedback", json={"query_id": unknown["query_id"], "dataset": "attendance"}, headers=company_headers
    )
    learned = _ask(client, company_headers, "ponches")
    for _ in range(3):
        useless = client.post(
            f"{URL}/feedback", json={"query_id": learned["query_id"], "helpful": False}, headers=company_headers
        )
        assert useless.json()["data"] == {"learned": [], "answer": None}
    assert _ask(client, company_headers, "ponches")["plan"] is None  # lo olvidó


def test_suggestions_come_from_what_the_company_asks(client, company_headers, company):
    for _ in range(2):
        _ask(client, company_headers, "empleados por departamento")
    catalog = client.get(f"{URL}/catalog", headers=company_headers).json()["data"]
    assert catalog["suggestions"][0] == "empleados por departamento"
    employees = next(d for d in catalog["datasets"] if d["code"] == "employees")
    department = next(c for c in employees["columns"] if c["code"] == "department")
    assert [o["label"] for o in department["options"]] == ["Producción", "Ventas"]


# ---------------------------------------------------------------- Excel


def test_export_to_excel_with_summary_and_safe_text(client, company_headers, company):
    plan = {"dataset": "employees", "columns": ["employee_number", "full_name", "department", "created_at"]}
    response = client.post(f"{URL}/export", json={"plan": plan, "question": "todos"}, headers=company_headers)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="reporte-empleados-' in response.headers["content-disposition"]
    book = load_workbook(BytesIO(response.content))
    data, summary = book["Datos"], book["Resumen"]
    rows = list(data.iter_rows(values_only=True))
    assert rows[0] == ("Número de empleado", "Nombre", "Departamento", "Alta") and len(rows) == 4
    malicious = next(c for c in data["B"] if c.value and str(c.value).startswith("="))
    assert malicious.data_type == "s"  # texto, nunca fórmula
    assert data["D2"].is_date
    assert ("Pregunta", "todos") in list(summary.iter_rows(values_only=True))


def test_export_of_groups_and_counts(client, company_headers, company):
    answer = _ask(client, company_headers, "empleados por departamento")
    exported = client.post(
        f"{URL}/export", json={"plan": answer["plan"], "query_id": answer["query_id"]}, headers=company_headers
    )
    rows = list(load_workbook(BytesIO(exported.content))["Datos"].iter_rows(values_only=True))
    assert rows == [("Departamento", "Cantidad"), ("Producción", 2), ("Ventas", 1)]
    with SessionLocal() as db:
        assert db.get(AssistantQuery, answer["query_id"]).exported is True  # type: ignore[union-attr]
    count = client.post(
        f"{URL}/export", json={"plan": {"dataset": "employees", "mode": "count"}}, headers=company_headers
    )
    assert list(load_workbook(BytesIO(count.content))["Datos"].iter_rows(values_only=True)) == [("Total",), (3,)]


# ---------------------------------------------------------------- aislamiento entre empresas


def test_never_reaches_another_company(client, company_headers, admin_headers, company):
    assert create_company(client, admin_headers).status_code == 201
    other = login(client, "admin@panificadora.com", "Empresa1234")
    stranger = _hire(client, other, "PAN-001", "Zoe", "Ajena", "zoe@panificadora.com")
    with SessionLocal() as db:
        _log(db, _company_of("admin@panificadora.com"), stranger, 10)
        db.commit()

    for dataset in ("employees", "attendance", "workdays", "departments", "validators", "devices", "enrollments"):
        exported = client.post(f"{URL}/export", json={"plan": {"dataset": dataset}}, headers=company_headers)
        values = {
            v for row in load_workbook(BytesIO(exported.content))["Datos"].iter_rows(values_only=True) for v in row
        }
        assert not any("Zoe" in str(v) or "PAN-001" in str(v) for v in values), dataset
    mine = _ask(client, other, "empleados")
    assert mine["preview"]["total"] == 1 and mine["preview"]["rows"][0][1] == "Zoe Ajena"
    # Ni sus preguntas ni sus reportes guardados.
    question = _ask(client, company_headers, "empleados")["query_id"]
    stolen = client.post(f"{URL}/feedback", json={"query_id": question, "helpful": True}, headers=other)
    assert stolen.status_code == 404 and stolen.json()["code"] == "REPORT_QUERY_NOT_FOUND"


@pytest.mark.parametrize("role_headers", ["admin_headers"])
def test_only_companies_use_the_assistant(client, request, role_headers, company):
    headers = request.getfixturevalue(role_headers)
    assert client.get(f"{URL}/catalog", headers=headers).status_code == 403
    employee = login(client, "ana@empresa.com", "Empleado123")
    assert client.post(f"{URL}/ask", json={"question": "empleados"}, headers=employee).status_code == 403


# ---------------------------------------------------------------- reportes guardados


def test_saved_reports(client, company_headers, admin_headers, company):
    answer = _ask(client, company_headers, "identificaciones de hoy por departamento")
    body = {"name": "Asistencia diaria", "plan": answer["plan"], "question": "hoy", "query_id": answer["query_id"]}
    created = client.post(f"{URL}/saved", json=body, headers=company_headers)
    assert created.status_code == 201 and created.json()["data"]["dataset_name"] == "Identificaciones"
    duplicate = client.post(f"{URL}/saved", json={**body, "name": "ASISTENCIA DIARIA"}, headers=company_headers)
    assert duplicate.status_code == 409 and duplicate.json()["code"] == "REPORT_NAME_TAKEN"

    report_id = created.json()["data"]["id"]
    run = client.post(f"{URL}/saved/{report_id}/run", headers=company_headers).json()["data"]
    assert run["plan"]["group_by"] == ["department"]
    listed = client.get(f"{URL}/saved", headers=company_headers).json()["data"]
    assert listed["total"] == 1 and listed["items"][0]["runs"] == 1

    assert create_company(client, admin_headers).status_code == 201
    other = login(client, "admin@panificadora.com", "Empresa1234")
    assert client.post(f"{URL}/saved/{report_id}/run", headers=other).json()["code"] == "SAVED_REPORT_NOT_FOUND"
    assert client.delete(f"{URL}/saved/{report_id}", headers=company_headers).status_code == 200
    assert client.get(f"{URL}/saved", headers=company_headers).json()["data"]["total"] == 0


# ---------------------------------------------------------------- planes inválidos


def _where(column: str, values: list, op: str = "in", dataset: str = "employees") -> dict:
    return {"dataset": dataset, "filters": [{"column": column, "op": op, "values": values}]}


@pytest.mark.parametrize(
    ("plan", "code"),
    [
        ({"dataset": "nomina"}, "REPORT_UNKNOWN_DATASET"),
        ({"dataset": "employees", "columns": ["salario"]}, "REPORT_UNKNOWN_COLUMN"),
        (_where("active", [True], "contains"), "REPORT_INVALID_FILTER"),
        (_where("active", []), "REPORT_INVALID_FILTER"),
        (_where("score", [1, 2], "gt", "attendance"), "REPORT_INVALID_FILTER"),
        (_where("active", ["si"]), "REPORT_INVALID_VALUE"),
        (_where("face_status", ["DROP"]), "REPORT_INVALID_VALUE"),
        (_where("full_name", [" "]), "REPORT_INVALID_VALUE"),
        (_where("birth_date", ["ayer"], "gt"), "REPORT_INVALID_VALUE"),
        (_where("score", [True], "gt", "attendance"), "REPORT_INVALID_VALUE"),
        ({"dataset": "employees", "mode": "groups", "group_by": ["email"]}, "REPORT_INVALID_GROUP"),
        ({"dataset": "employees", "mode": "groups"}, "REPORT_INVALID_GROUP"),
        ({"dataset": "employees", "mode": "count", "metric": "avg", "metric_column": "email"}, "REPORT_INVALID_METRIC"),
        ({"dataset": "employees", "sort": "salario"}, "REPORT_UNKNOWN_COLUMN"),
        (
            {"dataset": "attendance", "period": {"start": "2026-10-03T00:00:00Z", "end": "2026-10-01T00:00:00Z"}},
            "REPORT_INVALID_PERIOD",
        ),
    ],
)
def test_invalid_plans_are_rejected_with_their_code(client, company_headers, plan, code):
    for action in ("preview", "export"):  # también la descarga responde su error con el sobre
        response = client.post(f"{URL}/{action}", json={"plan": plan}, headers=company_headers)
        assert response.status_code == 422 and response.json()["code"] == code


def test_typed_filters_from_the_builder(client, company_headers, company):
    plan = {
        "dataset": "attendance",
        "filters": [
            {"column": "score", "op": "gte", "values": ["0.92"]},
            {"column": "occurred_at", "op": "gte", "values": ["2000-01-01T00:00:00"]},
            {"column": "employee", "op": "contains", "values": ["an"]},
            {"column": "method", "op": "not_in", "values": ["QR"]},
        ],
        "sort": "metric",
    }
    preview = client.post(f"{URL}/preview", json={"plan": plan}, headers=company_headers).json()["data"]
    assert preview["preview"]["total"] == 1 and preview["plan"]["sort"] == "metric"
    today = datetime.now(UTC).astimezone(business_day_start().tzinfo).date().isoformat()
    window = {"start": f"{today}T00:00:00", "end": None, "label": "hoy"}  # sin zona: la hora del negocio
    local = client.post(
        f"{URL}/preview", json={"plan": {"dataset": "attendance", "period": window}}, headers=company_headers
    )
    assert local.json()["data"]["preview"]["total"] == 4
    by_birth = {"dataset": "employees", "filters": [{"column": "birth_date", "op": "lte", "values": ["2000-01-01"]}]}
    assert (
        client.post(f"{URL}/preview", json={"plan": by_birth}, headers=company_headers).json()["data"]["preview"][
            "total"
        ]
        == 3
    )


def test_old_questions_are_purged(client, company_headers, company):
    _ask(client, company_headers, "empleados")
    with SessionLocal() as db:
        removed = maintenance_service.purge_expired(db, now=datetime.now(UTC) + timedelta(days=181))
        assert removed["preguntas al asistente de reportes"] >= 1


def test_time_groups_averages_and_long_previews(client, company_headers, company, monkeypatch):
    by_hour = _ask(client, company_headers, "identificaciones por hora de hoy")
    assert [row[0] for row in by_hour["preview"]["rows"]] == ["08:00", "09:00", "17:00"]
    by_month = _ask(client, company_headers, "identificaciones por mes")
    assert by_month["preview"]["rows"][-1][0].startswith(("octubre", "septiembre"))
    daily = _ask(client, company_headers, "identificaciones por día de esta semana")
    assert daily["plan"]["sort"] == "day" and len(daily["preview"]["rows"]) >= 1

    average = _ask(client, company_headers, "promedio de horas de hoy")
    assert average["answer"].startswith("Promedio de horas") and average["preview"]["columns"][0]["label"] == "Promedio"

    from app.core.config import settings

    monkeypatch.setattr(settings, "REPORT_PREVIEW_ROWS", 2)
    many = _ask(client, company_headers, "identificaciones")
    assert many["preview"]["truncated"] is True and "Te muestro 2" in many["answer"]


def test_feedback_edge_cases(client, company_headers, company):
    help_ = _ask(client, company_headers, "hola")
    useful = client.post(
        f"{URL}/feedback", json={"query_id": help_["query_id"], "helpful": True}, headers=company_headers
    )
    assert useful.status_code == 200 and useful.json()["data"]["learned"] == []
    bad = client.post(
        f"{URL}/feedback", json={"query_id": help_["query_id"], "dataset": "nomina"}, headers=company_headers
    )
    assert bad.status_code == 422 and bad.json()["code"] == "REPORT_UNKNOWN_DATASET"

    # Entendió «ponches» como empleados (por un dato de la conversación) y la persona lo corrige.
    wrong = _ask(client, company_headers, "ponches de los empleados")
    assert wrong["plan"]["dataset"] == "employees"
    fixed = client.post(
        f"{URL}/feedback", json={"query_id": wrong["query_id"], "dataset": "attendance"}, headers=company_headers
    ).json()["data"]
    assert fixed["learned"] == ["«ponche» → Identificaciones"]

    # Guardar sin respuesta del asistente, o con una que no tenía datos, no enseña nada.
    plan = {"dataset": "employees"}
    plain = client.post(f"{URL}/saved", json={"name": "Plantilla", "plan": plan}, headers=company_headers)
    assert plain.status_code == 201
    from_help = {"name": "Ayuda", "plan": plan, "query_id": help_["query_id"]}
    assert client.post(f"{URL}/saved", json=from_help, headers=company_headers).status_code == 201
