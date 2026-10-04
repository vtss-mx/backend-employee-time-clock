"""El intérprete del asistente: español real (periodos, errores de dedo, conversación) → plan.

Unitarias (sin HTTP): el intérprete, los periodos, el lenguaje, el narrador y el archivo de Excel.
El catálogo de la BD (motivos, métodos...) sí se usa: lo siembra el fixture de cada prueba.
"""

from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.schemas.report import ReportFilter, ReportPlan
from app.services.catalog_service import get_catalogs
from app.services.reporting import excel, language, narrator
from app.services.reporting.datasets import BY_CODE
from app.services.reporting.interpreter import Interpreter, Lexicon
from app.services.reporting.periods import Period, find_period

TODAY = date(2026, 10, 3)  # sábado


def _interpret(question: str, context: ReportPlan | None = None, **lexicon):
    words = {"departments": ["Producción", "Ventas"], "validators": ["Caseta Norte"], "learned": {}}
    words |= {"employees": [("Ana López García", "EMP-001"), ("Beto Ruiz", "EMP-002")]}
    words |= lexicon
    return Interpreter(Lexicon(catalogs=get_catalogs(), **words), TODAY).interpret(question, context)


def _filters(result) -> dict:
    return {f.column: (f.op, f.values) for f in result.plan.filters}


# ---------------------------------------------------------------- periodos


@pytest.mark.parametrize(
    ("text", "first", "last", "label"),
    [
        ("hace 2 semanas", date(2026, 9, 14), date(2026, 9, 21), "del 14 al 20 de septiembre de 2026"),
        ("hace un mes", date(2026, 9, 1), date(2026, 10, 1), "septiembre de 2026"),
        ("hace 3 dias", date(2026, 9, 30), date(2026, 10, 1), "30 de septiembre de 2026"),
        (
            "del 28 de diciembre de 2025 al 3 de enero",
            date(2025, 12, 28),
            date(2026, 1, 4),
            "del 28 de diciembre de 2025 al 3 de enero de 2026",
        ),
        ("entre 1/3/26 y 15/3/26", date(2026, 3, 1), date(2026, 3, 16), "del 1 al 15 de marzo de 2026"),
        ("desde el 15/09/2026", date(2026, 9, 15), date(2026, 10, 4), "desde el 15 de septiembre de 2026"),
        ("a partir del 1 de agosto", date(2026, 8, 1), date(2026, 10, 4), "desde el 1 de agosto de 2026"),
        ("el ultimo ano", date(2025, 10, 4), date(2026, 10, 4), "último año"),
        ("los ultimos 2 meses", date(2026, 8, 5), date(2026, 10, 4), "últimos 60 días".replace("60 días", "2 meses")),
        ("el miercoles", date(2026, 9, 30), date(2026, 10, 1), "el miércoles 30 de septiembre de 2026"),
        ("el sabado", date(2026, 10, 3), date(2026, 10, 4), "el sábado 3 de octubre de 2026"),
        ("durante 2024", date(2024, 1, 1), date(2025, 1, 1), "2024"),
        ("el 3 de octubre", date(2026, 10, 3), date(2026, 10, 4), "3 de octubre de 2026"),
        ("en marzo", date(2026, 3, 1), date(2026, 4, 1), "marzo de 2026"),
        ("en noviembre", date(2025, 11, 1), date(2025, 12, 1), "noviembre de 2025"),
        ("diciembre de 2024", date(2024, 12, 1), date(2025, 1, 1), "diciembre de 2024"),
        ("del 3 al 3 de marzo", date(2026, 3, 3), date(2026, 3, 4), "3 de marzo de 2026"),
        ("el sabado pasado", date(2026, 9, 26), date(2026, 9, 27), "el sábado 26 de septiembre de 2026"),
    ],
)
def test_periods_in_spanish(text, first, last, label):
    found = find_period(language.normalize(text), TODAY)
    assert found is not None
    period = found.period
    assert period.start is not None
    assert (period.start.date(), period.end.date(), period.label) == (first, last, label)


@pytest.mark.parametrize(
    "text",
    [
        "el 31/02/2026",
        "del 15 al 3 de marzo",
        "del 31/02/2026 al 01/03/2026",
        "del 10/03 al 01/03",
        "desde el 31/02/2026",
    ],
)
def test_impossible_dates_are_not_periods(text):
    found = find_period(language.normalize(text), TODAY)
    assert found is None or found.period.label != text


def test_until_a_date_has_no_start():
    found = find_period("hasta el 15/09/2026", TODAY)
    assert (
        found is not None and found.period.start is None and found.period.label == "hasta el 15 de septiembre de 2026"
    )
    impossible = find_period("el 31 de febrero", TODAY)  # no existe: se busca otro periodo (febrero)
    assert impossible is not None and impossible.period.label == "febrero de 2026"


def test_previous_period_has_the_same_length():
    found = find_period("este mes", TODAY)
    assert found is not None
    previous = found.period.previous()
    assert previous is not None and previous.end == found.period.start
    assert Period(None, found.period.end, "hasta hoy").previous() is None


# ---------------------------------------------------------------- lenguaje


def test_language_helpers():
    assert language.stem("qr") == "qr" and language.stem("a1b2s") == "a1b2s" and language.stem("veces") == "vez"
    assert language.similar("rostro", "rostro")
    assert not language.similar("validador", "validadoresxyz")  # largo muy distinto
    assert language.find_phrase(["hola"], ()) == -1 and language.find_phrase(["a"], ("a", "b")) == -1
    assert language.number_at("diez") == 10 and language.number_at("12") == 12 and language.number_at("x") is None


# ---------------------------------------------------------------- intérprete


def test_entities_by_number_and_by_first_name_plus_surname():
    by_number = _interpret("identificaciones de EMP-002")
    assert _filters(by_number) == {"employee": ("in", ["Beto Ruiz"])}
    partial = _interpret("identificaciones de Ana López")  # sin el segundo apellido
    assert _filters(partial) == {"employee": ("in", ["Ana López García"])}
    unrelated = _interpret("departamentos de Caseta Norte")
    assert any("no se puede filtrar por Caseta Norte" in note for note in unrelated.notes)


def test_defaults_when_no_dataset_is_named():
    assert _interpret("¿cuántos hay en Ventas?").plan.dataset == "employees"
    assert _interpret("retrasos de ayer").plan.dataset == "workdays"
    assert _interpret("¿qué pasó ayer?").plan.dataset == "attendance"
    weak = _interpret("nombre")  # una sola columna: se intenta, con poca confianza
    assert weak.plan.dataset == "employees" and weak.confidence < 1


def test_negations_comparisons_and_times_of_day():
    failed = _interpret("identificaciones no exitosas")
    assert _filters(failed) == {"success": ("in", [False])}
    big = _interpret("departamentos con más de 5 empleados")
    assert _filters(big) == {"employees": ("gt", [5.0])}
    unknown_number = _interpret("departamentos con más de 5 sillas")
    assert unknown_number.plan.filters == []
    bare = _interpret("departamentos con más de 5")
    assert bare.plan.filters == []
    both = _interpret("identificaciones fallidas y rechazadas")  # dos palabras, el mismo valor
    assert _filters(both) == {"success": ("in", [False])}
    leaving = _interpret("¿quién salió antes de las 5 pm?")
    assert _filters(leaving) == {"exit_hour": ("lt", [17.0])} or leaving.plan.dataset == "workdays"
    by_hour = _interpret("identificaciones después de las 6:30 de la tarde")
    assert _filters(by_hour) == {"hour": ("gt", [18.5])}


def test_groups_top_metrics_and_order():
    daily = _interpret("identificaciones por día de la semana pasada")
    assert daily.plan.group_by == ["day"] and daily.plan.sort == "day" and daily.plan.descending is False
    why = _interpret("identificaciones fallidas, ¿por qué?, por favor")
    assert why.plan.group_by == ["reason"]
    twice = _interpret("identificaciones por departamento por departamento")
    assert twice.plan.group_by == ["department"]
    top = _interpret("los 5 empleados con menos identificaciones")
    assert top.plan.limit == 5 and top.plan.descending is False and top.plan.group_by == ["employee"]
    average = _interpret("promedio de horas de hoy")
    assert (average.plan.mode, average.plan.metric, average.plan.metric_column) == ("count", "avg", "hours")
    confidence = _interpret("promedio de confianza por validador")
    assert (confidence.plan.metric, confidence.plan.metric_column) == ("avg", "score")
    devices = _interpret("promedio de dispositivos de los validadores por ciudad")
    assert (devices.plan.dataset, devices.plan.metric_column) == ("validators", "devices")
    unknown_group = _interpret("identificaciones por zzz")
    assert unknown_group.plan.mode == "rows"
    no_metric = _interpret("promedio de empleados por departamento")
    assert no_metric.plan.metric == "count"
    recent = _interpret("los registros faciales más recientes")
    assert (recent.plan.sort, recent.plan.descending) == ("submitted_at", True)
    oldest = _interpret("empleados más antiguos")
    assert (oldest.plan.sort, oldest.plan.descending) == ("created_at", False)
    contact = _interpret("empleados con su correo y teléfono y rfc")
    assert {"phone", "rfc"} <= set(contact.plan.columns)


def test_conversation_follow_ups():
    first = _interpret("identificaciones fallidas de hoy").plan
    cleared = _interpret("y ahora todas, sin filtros", first)
    assert cleared.plan.filters == [] and cleared.plan.period == first.period
    detail = _interpret("¿cuáles son?", first.model_copy(update={"mode": "count"}))
    assert detail.plan.mode == "rows"
    forged = first.model_copy(update={"filters": [ReportFilter(column="salario", values=["1"])]})
    assert _interpret("ahora de ayer", forged).plan.period.label.startswith("ayer")


def test_periods_on_data_without_dates_are_explained():
    learning = _interpret("aprendizaje facial de ayer")
    assert learning.plan.period is None and any("no tiene fecha" in note for note in learning.notes)


def test_learned_words_count_only_for_their_dataset():
    learned = {"ponche": {"attendance": 3}, "otro": {"employees": 1}}
    result = _interpret("ponches de hoy", learned=learned)
    assert result.plan.dataset == "attendance" and result.learned == ["ponche"]


# ---------------------------------------------------------------- narrador y Excel


def test_narration_details():
    attendance = BY_CODE["attendance"]
    assert narrator.amount(1234) == "1,234" and narrator.amount(7.456) == "7.46"
    same, _ = narrator.narrate(
        attendance, "count", narrator.Outcome(total=5, groups=[], previous=5), period=None, conditions=[]
    )
    assert same == "Hay 5 identificaciones."
    _, highlights = narrator.narrate(
        attendance, "count", narrator.Outcome(total=5, groups=[], previous=5), period="hoy", conditions=[]
    )
    assert highlights == ["Igual que en el periodo anterior (5)."]
    _, new = narrator.narrate(
        attendance, "count", narrator.Outcome(total=2, groups=[], previous=0), period=None, conditions=[]
    )
    assert new == ["No hubo ninguno en el periodo anterior."]
    _, down = narrator.narrate(
        attendance, "count", narrator.Outcome(total=1, groups=[], previous=4), period=None, conditions=["x"]
    )
    assert down == ["75 % menos que en el periodo anterior (4)."]
    one_group = narrator.Outcome(total=1, groups=[("Ana", 3.0)])
    _, single = narrator.narrate(attendance, "groups", one_group, period=None, conditions=[], group_name="Empleado")
    assert not any("Al final" in h for h in single)
    metric = narrator.Outcome(total=1, groups=[("Ana", 7.5)])
    _, measured = narrator.narrate(
        attendance, "groups", metric, period=None, conditions=[], metric_name="Promedio de horas", group_name="Empleado"
    )
    assert measured == ["Ana: 7.5."]
    empty_company, _ = narrator.narrate(
        attendance, "rows", narrator.Outcome(total=0, groups=[]), period=None, conditions=[]
    )
    assert empty_company == "No hay identificaciones."
    nothing, _ = narrator.overview([(attendance, 0)], [])
    assert nothing.startswith("Aún no hay datos")


def test_excel_types_and_an_empty_report():
    columns = [
        excel.SheetColumn("Día", "date"),
        excel.SheetColumn("Veces", "number"),
        excel.SheetColumn("Nota", "text"),
    ]
    summary = excel.Summary("Reporte: prueba", "Empresa", None, ["Periodo: hoy"], datetime(2026, 10, 3, 9), limit=1)
    content = excel.build_workbook(columns, [[date(2026, 10, 3), 7, "a\x01b"]], summary)
    book = load_workbook(BytesIO(content))
    data, info = book["Datos"], book["Resumen"]
    assert data["A2"].is_date and data["B2"].number_format == "#,##0" and data["C2"].value == "ab"
    assert any(row[0] == "Aviso" for row in info.iter_rows(values_only=True))  # se cortó en 1 fila
    blank = load_workbook(BytesIO(excel.build_workbook([], [], summary)))
    assert blank["Datos"].max_row == 1
