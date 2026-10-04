"""Archivo de Excel de un reporte: hoja «Datos» y hoja «Resumen» (qué se consultó y cuándo).

- Modo de solo escritura de openpyxl: las filas se escriben conforme llegan de la BD; un reporte de
  100 000 filas no se carga entero en memoria.
- Fechas en la hora del negocio y con formato de fecha de Excel (se pueden filtrar y ordenar).
- Seguro: un texto que empieza con «=» (p. ej. un nombre capturado con mala intención) se guarda
  como texto, nunca como fórmula; los caracteres que Excel no acepta se quitan.
"""

from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._write_only import WriteOnlyWorksheet

from app.core.clock import as_utc
from app.core.config import settings

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_TITLE_FONT = Font(bold=True, size=14)
_LABEL_FONT = Font(bold=True)
_WIDTHS = {"text": 28, "category": 22, "datetime": 19, "date": 12, "number": 14, "bool": 12}
_FORMATS = {"datetime": "dd/mm/yyyy hh:mm", "date": "dd/mm/yyyy"}


@dataclass(frozen=True)
class SheetColumn:
    label: str
    kind: str


@dataclass(frozen=True)
class Summary:
    title: str
    company: str
    question: str | None
    understood: Sequence[str]
    generated_at: datetime
    #: Tope de filas si se cortó el reporte.
    limit: int


def _text(value: str) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", value)


def _cell(sheet: WriteOnlyWorksheet, value: Any, kind: str) -> WriteOnlyCell:
    """Celda con su tipo: fecha local sin zona (Excel no las admite), número, o texto (nunca fórmula)."""
    if isinstance(value, datetime):
        value = as_utc(value).astimezone(ZoneInfo(settings.APP_TIMEZONE)).replace(tzinfo=None)
    elif isinstance(value, str):
        value = _text(value)
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        cell.data_type = "s"  # «=HYPERLINK(...)» queda como texto
    elif isinstance(value, (datetime, date)):
        cell.number_format = _FORMATS["datetime" if isinstance(value, datetime) else "date"]
    elif isinstance(value, float):
        cell.number_format = "#,##0.00"
    elif kind == "number" and isinstance(value, int):
        cell.number_format = "#,##0"
    return cell


def _header(sheet: WriteOnlyWorksheet, label: str) -> WriteOnlyCell:
    cell = WriteOnlyCell(sheet, value=_text(label))
    cell.data_type = "s"
    cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
    cell.alignment = Alignment(vertical="center")
    return cell


def _summary(book: Workbook, summary: Summary, rows: int) -> None:
    sheet = book.create_sheet("Resumen")
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 90
    title = WriteOnlyCell(sheet, value=_text(summary.title))
    title.font, title.data_type = _TITLE_FONT, "s"
    sheet.append([title])
    generated = summary.generated_at.strftime("%d/%m/%Y %H:%M")
    lines: list[tuple[str, str]] = [
        ("Empresa", summary.company),
        ("Generado", generated),
        ("Filas", f"{rows:,}"),
    ]
    if summary.question:
        lines.insert(1, ("Pregunta", summary.question))
    lines += [("Consulta", item) for item in summary.understood]
    if rows >= summary.limit:
        lines.append(("Aviso", f"El reporte se cortó en {summary.limit:,} filas: agrega filtros o un periodo."))
    for label, value in lines:
        name = WriteOnlyCell(sheet, value=label)
        name.font = _LABEL_FONT
        sheet.append([name, _cell(sheet, value, "text")])


def build_workbook(columns: Sequence[SheetColumn], rows: Iterable[Sequence[Any]], summary: Summary) -> bytes:
    """El archivo .xlsx completo. `rows` puede ser un generador (se consume una vez)."""
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Datos")
    for index, column in enumerate(columns, start=1):
        width = max(_WIDTHS.get(column.kind, 18), min(len(column.label) + 4, 40))
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.append([_header(sheet, column.label) for column in columns])
    written = 0
    for row in rows:
        sheet.append([_cell(sheet, value, column.kind) for value, column in zip(row, columns, strict=True)])
        written += 1
    if columns:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{written + 1}"
    _summary(book, summary, written)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
